from flask import Blueprint, request, jsonify, g, send_file
from flask_jwt_extended import get_jwt_identity
from datetime import datetime, timedelta, date
from bson import ObjectId
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import csv, io

from auth_utils import tenant_scoped, require_role
from tenant_scope import get_db

employees_bp = Blueprint('employees', __name__)

DELETE_ROLES = {'admin', 'hr_head'}

BULK_UPLOAD_COLUMNS = [
    ('name',             'Name', True),
    ('designation',      'Designation', True),
    ('department',       'Department', False),
    ('email',            'Email', False),
    ('phone',            'Phone', False),
    ('address',          'Address', False),
    ('joining_date',     'Joining Date (DD-MM-YYYY)', False),
    ('date_of_birth',    'Date of Birth (DD-MM-YYYY)', False),
    ('ctc',              'CTC', False),
    ('basic',            'Basic', False),
    ('hra',              'HRA', False),
    ('da',               'DA', False),
    ('allowances',       'Allowances', False),
    ('probation_period', 'Probation Period (months)', False),
    ('notice_period',    'Notice Period (days)', False),
]

def serialize(emp):
    emp['_id'] = str(emp['_id'])
    return emp


def _next_emp_id(db):
    counter = db.counters.find_one_and_update(
        {'name': 'employee_id'},
        {'$inc': {'seq': 1}},
        upsert=True,
        return_document=True,
    )
    return f"EMP{str(counter['seq']).zfill(3)}"


@employees_bp.route('/', methods=['GET'])
@tenant_scoped
def list_employees():
    db     = get_db()
    status = request.args.get('status')
    query  = {}
    if status:
        query['status'] = status
    
    all_emps = list(db.employees.find(query).sort('created_at', -1))
    
    # Only show employees whose offer letter has reached id_created or beyond
    VISIBLE_LETTER_STATUSES = {'id_created', 'joined'}
    
    visible = []
    for emp in all_emps:
        emp_id = str(emp['_id'])
        # Non-active statuses (exiting, exited, inactive) always show
        if emp.get('status') != 'active':
            visible.append(serialize(emp))
            continue
        # Active: only show if they have a letter at id_created or joined
        letter = db.letters.find_one({
            'employee_id': emp_id,
            'status': {'$in': list(VISIBLE_LETTER_STATUSES)}
        })
        if letter:
            visible.append(serialize(emp))
    
    return jsonify(visible)


@employees_bp.route('/me/step1', methods=['POST'])
@tenant_scoped
def save_step1():
    """
    Employee submits Step 1 personal info form.
    Saves full form under step1_data on the employee record,
    and syncs key fields to root level for the employee card.
    """
    db   = get_db()
    user = g.caller

    emp_ref = user.get('employee_ref')
    if not emp_ref:
        return jsonify({'error': 'No employee record linked to this account'}), 400

    data = request.json or {}
    if not data:
        return jsonify({'error': 'No data provided'}), 400

    data.pop('_id', None)

    # Frontend now sends structured address fields (address_line1/2, city,
    # state, postal_code, country) plus repeatable education/work_experience/
    # dependents arrays. Build a single-line address string for backward
    # compatibility with letter generation ({{address}} placeholder) and
    # older HR views that expect a flat 'address' string.
    address = data.get('address') or ', '.join(filter(None, [
        data.get('address_line1'), data.get('address_line2'),
        data.get('city'), data.get('state'),
        data.get('postal_code'), data.get('country'),
    ]))

    db.employees.update_one(
        {'_id': ObjectId(emp_ref)},
        {'$set': {
            'step1_data':    data,
            'phone':         data.get('phone', ''),
            'date_of_birth': data.get('dob', ''),
            'address':       address,
            'pan_number':    data.get('pan_number', ''),
            'bank_account':  data.get('account_number', ''),
            'bank_ifsc':     data.get('ifsc_code', ''),
            'updated_at':    datetime.utcnow(),
        }}
    )
    return jsonify({'message': 'Personal information saved successfully.'}), 200



@employees_bp.route('/bulk-upload/template', methods=['GET'])
@require_role('admin', 'hr', 'hr_head')
def bulk_upload_template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Employees'
    headers = [label for _, label, _ in BULK_UPLOAD_COLUMNS]
    ws.append(headers)
    for col_idx, (_, _, required) in enumerate(BULK_UPLOAD_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = cell.font.copy(bold=True)
    ws.append([
        'Jane Doe', 'Software Engineer', 'Engineering', 'jane.doe@example.com',
        '9876543210', '123 Main St, City', '01-06-2026', '15-03-1995',
        800000, 40000, 15000, 5000, 10000, 6, 30,
    ])
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='employee_bulk_upload_template.xlsx',
    )


@employees_bp.route('/bulk-upload', methods=['POST'])
@require_role('admin', 'hr', 'hr_head')
def bulk_upload():
    db   = get_db()
    file = request.files.get('file')
    if not file:
        return jsonify({'error': 'No file provided'}), 400

    filename = (file.filename or '').lower()
    rows = []  # list of dicts keyed by our internal field names

    if filename.endswith('.xlsx'):
        try:
            wb = load_workbook(file.stream, data_only=True)
        except Exception:
            return jsonify({'error': 'Could not read the Excel file. Please use the downloadable template.'}), 400
        ws = wb.active
        header_row = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
        label_to_field = {label: field for field, label, _ in BULK_UPLOAD_COLUMNS}
        field_by_col = [label_to_field.get(h) for h in header_row]
        for excel_row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None or str(v).strip() == '' for v in excel_row):
                continue
            row = {}
            for col_idx, value in enumerate(excel_row):
                field = field_by_col[col_idx] if col_idx < len(field_by_col) else None
                if field and value is not None:
                    row[field] = str(value).strip() if not isinstance(value, (int, float)) else value
            rows.append(row)
    elif filename.endswith('.csv'):
        content = file.stream.read().decode('utf-8-sig', errors='replace')
        stream  = io.StringIO(content)
        reader  = csv.DictReader(stream)
        rows = [dict(r) for r in reader]
    else:
        return jsonify({'error': 'Unsupported file type. Please upload the .xlsx template or a .csv file.'}), 400

    if not rows:
        return jsonify({'error': 'No employee rows found in the file.'}), 400

    imported, errors = [], []
    for i, row in enumerate(rows, start=2):  # row 1 is the header
        name = str(row.get('name', '')).strip()
        designation = str(row.get('designation', '')).strip()
        if not name or not designation:
            errors.append({'row': i, 'error': 'Name and Designation are required.'})
            continue

        doc = {k: v for k, v in row.items() if v not in (None, '')}
        doc['name'] = name
        doc['designation'] = designation
        doc['employee_id'] = doc.get('employee_id') or _next_emp_id(db)
        doc['created_at']  = datetime.utcnow()
        doc['status']      = doc.get('status', 'active')
        imported.append(doc)

    if imported:
        db.employees.insert_many(imported)

    return jsonify({
        'message': f'{len(imported)} employee(s) imported' + (f', {len(errors)} row(s) skipped' if errors else ''),
        'imported': len(imported),
        'errors': errors,
    })


@employees_bp.route('/<emp_id>', methods=['GET'])
@tenant_scoped
def get_employee(emp_id):
    db  = get_db()
    emp = db.employees.find_one({'_id': ObjectId(emp_id)})
    if not emp:
        return jsonify({'error': 'Not found'}), 404
    return jsonify(serialize(emp))


@employees_bp.route('/', methods=['POST'])
@tenant_scoped
def create_employee():
    db   = get_db()
    data = request.json or {}
    if not data.get('name') or not data.get('designation'):
        return jsonify({'error': 'name and designation are required'}), 400
    if not data.get('employee_id'):
        data['employee_id'] = _next_emp_id(db)
    data['created_at'] = datetime.utcnow()
    data['status']     = data.get('status', 'active')
    result = db.employees.insert_one(data)
    return jsonify({'id': str(result.inserted_id), 'employee_id': data['employee_id']}), 201


@employees_bp.route('/<emp_id>', methods=['PUT'])
@require_role('admin')
def update_employee(emp_id):
    db     = get_db()
    caller = g.caller

    data = request.json or {}
    data.pop('_id', None)
    data['updated_at'] = datetime.utcnow()
    db.employees.update_one({'_id': ObjectId(emp_id)}, {'$set': data})
    return jsonify({'message': 'Updated'})


@employees_bp.route('/<emp_id>', methods=['DELETE'])
@require_role(*DELETE_ROLES)
def delete_employee(emp_id):
    db     = get_db()
    uid    = get_jwt_identity()
    caller = g.caller

    emp = db.employees.find_one({'_id': ObjectId(emp_id)})
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404

    status = emp.get('status', 'active')

    if status in ('notice_period', 'clearance_pending', 'clearance_complete'):
        return jsonify({
            'error': (
                f"Cannot delete employee with status '{status}'. "
                "This employee is mid-exit-process. Complete or cancel the exit workflow first."
            )
        }), 400

    if status == 'active':
        linked_letters = db.letters.count_documents({'employee_id': emp_id})
        if linked_letters > 0:
            return jsonify({
                'error': (
                    f"Cannot delete: {linked_letters} offer letter(s) exist for this employee. "
                    "Delete or archive the letters first, or mark the employee as exited."
                )
            }), 400

    db.employees.delete_one({'_id': ObjectId(emp_id)})
    return jsonify({
        'message': f"Employee '{emp.get('name')}' deleted by {caller.get('name', uid)} ({caller.get('role')}) at {datetime.utcnow().isoformat()}"
    })


@employees_bp.route('/<emp_id>/deactivate', methods=['POST'])
@require_role('admin', 'hr_head')
def deactivate_employee(emp_id):
    db     = get_db()
    uid    = get_jwt_identity()
    caller = g.caller

    emp = db.employees.find_one({'_id': ObjectId(emp_id)})
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404

    if emp.get('status') != 'exited':
        return jsonify({'error': 'Only exited employees can be deactivated'}), 400

    letter = db.letters.find_one({'employee_id': emp_id, 'letter_type': 'relieving'})
    if not letter:
        return jsonify({'error': 'Relieving letter must be issued before deactivating the account'}), 400

    db.employees.update_one(
        {'_id': ObjectId(emp_id)},
        {'$set': {
            'status':         'inactive',
            'deactivated_by': uid,
            'deactivated_at': datetime.utcnow().isoformat(),
            'updated_at':     datetime.utcnow(),
        }}
    )
    db.users.update_one(
        {'employee_ref': emp_id},
        {'$set': {'is_active': False, 'updated_at': datetime.utcnow()}}
    )
    return jsonify({'message': 'Employee deactivated. Login access revoked.'})

@employees_bp.route('/<emp_id>/activate', methods=['POST'])
@require_role('admin', 'hr_head')
def activate_employee(emp_id):
    db     = get_db()
    uid    = get_jwt_identity()
    caller = g.caller

    emp = db.employees.find_one({'_id': ObjectId(emp_id)})
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404

    if emp.get('status') != 'inactive':
        return jsonify({'error': 'Only inactive employees can be activated'}), 400

    db.employees.update_one(
        {'_id': ObjectId(emp_id)},
        {'$set': {
            'status':       'active',
            'activated_by': uid,
            'activated_at': datetime.utcnow().isoformat(),
            'updated_at':   datetime.utcnow(),
        }}
    )
    db.users.update_one(
        {'employee_ref': emp_id},
        {'$set': {'is_active': True, 'updated_at': datetime.utcnow()}}
    )
    return jsonify({'message': 'Employee account activated. Login access restored.'})


# ── Exit workflow routes ──────────────────────────────────────────────────────

@employees_bp.route('/<emp_id>/exit', methods=['POST'])
@tenant_scoped
def record_exit(emp_id):
    db   = get_db()
    data = request.json or {}
    emp  = db.employees.find_one({'_id': ObjectId(emp_id)})
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404
    if emp.get('status') not in ('active',):
        return jsonify({'error': f"Cannot record resignation: employee status is '{emp.get('status')}'"}), 400

    resignation_date = data.get('resignation_date')
    if not resignation_date:
        return jsonify({'error': 'resignation_date is required'}), 400

    lwd = data.get('last_working_day')
    if not lwd:
        notice_days = int(emp.get('notice_period', 60))
        lwd = (datetime.strptime(resignation_date, '%Y-%m-%d') + timedelta(days=notice_days)).strftime('%Y-%m-%d')

    db.employees.update_one(
        {'_id': ObjectId(emp_id)},
        {'$set': {
            'status':           'notice_period',
            'resignation_date': resignation_date,
            'last_working_day': lwd,
            'exit_reason':      data.get('exit_reason', ''),
            'clearances': {
                'it_assets': False, 'finance': False, 'admin': False,
                'hr_docs': False, 'access_cards': False,
            },
            'updated_at': datetime.utcnow(),
        }}
    )
    return jsonify({'message': 'Resignation recorded', 'last_working_day': lwd}), 200


@employees_bp.route('/<emp_id>/clearance', methods=['POST'])
@tenant_scoped
def update_clearance(emp_id):
    db   = get_db()
    data = request.json or {}
    emp  = db.employees.find_one({'_id': ObjectId(emp_id)})
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404

    VALID_KEYS  = {'it_assets', 'finance', 'admin', 'hr_docs', 'access_cards'}
    updates     = {k: bool(v) for k, v in data.items() if k in VALID_KEYS}
    if not updates:
        return jsonify({'error': 'No valid clearance keys provided'}), 400

    current     = emp.get('clearances', {})
    merged      = {**current, **updates}
    all_cleared = all(merged.get(k, False) for k in VALID_KEYS)
    new_status  = 'clearance_complete' if all_cleared else 'clearance_pending'

    db.employees.update_one(
        {'_id': ObjectId(emp_id)},
        {'$set': {
            'clearances': merged,
            'status':     new_status,
            'updated_at': datetime.utcnow(),
        }}
    )
    return jsonify({'message': 'Clearance updated', 'all_cleared': all_cleared, 'status': new_status})


# ─────────────────────────────────────────────────────────────────────────────
# Birthdays & work anniversaries — in-app widget data. Mirrors the date
# matching scheduler.py uses for its email job, but computed live here for
# display rather than as a side-effecting daily job.
# ─────────────────────────────────────────────────────────────────────────────

def _parse_wish_date(d_str):
    if not d_str or not isinstance(d_str, str):
        return None
    d_str = d_str.strip()
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(d_str, fmt).date()
        except ValueError:
            continue
    return None


def _next_occurrence(month, day, today):
    """Days until the next time (month, day) occurs, treating today as 0."""
    try:
        this_year = date(today.year, month, day)
    except ValueError:
        this_year = date(today.year, month, min(day, 28))  # Feb 29 in a non-leap year
    if this_year >= today:
        return (this_year - today).days
    try:
        next_year = date(today.year + 1, month, day)
    except ValueError:
        next_year = date(today.year + 1, month, min(day, 28))
    return (next_year - today).days


@employees_bp.route('/wishes', methods=['GET'])
@tenant_scoped
def wishes():
    db = get_db()
    try:
        days_ahead = int(request.args.get('days', 7))
    except ValueError:
        days_ahead = 7
    days_ahead = max(0, min(days_ahead, 60))

    today = date.today()

    users_by_empref = {
        u['employee_ref']: u for u in db.users.find({
            'employee_ref': {'$exists': True, '$ne': ''},
            'is_active': {'$ne': False},
        })
    }
    employees = list(db.employees.find({'status': 'active'}))

    items = []
    for emp in employees:
        emp_id = str(emp['_id'])
        user = users_by_empref.get(emp_id, {})
        name = user.get('name') or emp.get('name', 'Team Member')

        bday_str = user.get('birthday') or emp.get('birthday') or emp.get('date_of_birth', '')
        bday = _parse_wish_date(bday_str)
        if bday:
            days_until = _next_occurrence(bday.month, bday.day, today)
            if days_until <= days_ahead:
                items.append({
                    'employee_id': emp_id, 'name': name, 'type': 'birthday',
                    'days_until': days_until, 'month': bday.month, 'day': bday.day,
                })

        join_str = emp.get('joining_date', '')
        join_dt = _parse_wish_date(join_str)
        if join_dt:
            years = today.year - join_dt.year
            days_until = _next_occurrence(join_dt.month, join_dt.day, today)
            occurrence_year = today.year if days_until == 0 or (today + timedelta(days=days_until)).year == today.year else today.year + 1
            years_completing = occurrence_year - join_dt.year
            if years_completing >= 1 and days_until <= days_ahead:
                items.append({
                    'employee_id': emp_id, 'name': name, 'type': 'anniversary',
                    'days_until': days_until, 'years': years_completing,
                    'month': join_dt.month, 'day': join_dt.day,
                })

    items.sort(key=lambda i: i['days_until'])
    return jsonify({
        'today': [i for i in items if i['days_until'] == 0],
        'upcoming': [i for i in items if i['days_until'] > 0],
    })