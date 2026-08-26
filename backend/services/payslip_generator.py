"""
payslip_generator.py — builds a payslip DOCX (then PDF, via the same
LibreOffice conversion letters already use) directly with python-docx.

Unlike offer letters, payslips don't need a hand-authored template with
{{placeholders}} — the layout is a fixed earnings/deductions statement, so
it's built programmatically here. generate_letter_pdf() (letter_generator.py)
is reused as-is for the DOCX -> PDF step so both document types go through
the exact same, already-proven conversion path.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MONTH_NAMES = ['', 'January', 'February', 'March', 'April', 'May', 'June',
               'July', 'August', 'September', 'October', 'November', 'December']

ACCENT = RGBColor(0x2A, 0x78, 0xD6)
MUTED = RGBColor(0x5B, 0x66, 0x77)
BORDER_GRAY = 'D0D5DD'


def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _set_cell_borders(cell, color=BORDER_GRAY):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def _row(table, label, value, bold=False, shade=None):
    cells = table.add_row().cells
    cells[0].text = label
    cells[1].text = value
    for c in cells:
        _set_cell_borders(c)
        if shade:
            _shade_cell(c, shade)
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10.5)
                r.font.bold = bold
    cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return cells


def _inr(n):
    return f"Rs. {int(round(n or 0)):,}"


def generate_payslip_docx(payslip, output_path):
    """
    payslip: the serialized payslip dict (already enriched with
    employee_name/employee_code/designation/department by the route).
    """
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.75)
    section.top_margin = section.bottom_margin = Inches(0.6)

    company = os.getenv('COMPANY_NAME', 'Infopace Management Pvt Ltd')
    month_name = MONTH_NAMES[int(payslip.get('month', 0))] if payslip.get('month') else ''

    # ── Header ──────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(company)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f"Payslip for {month_name} {payslip.get('year', '')}")
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED

    status = (payslip.get('status') or 'draft').upper()
    watermark = doc.add_paragraph()
    watermark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = watermark.add_run(f"[ {status} — NOT A FINAL DOCUMENT UNTIL RELEASED ]" if status != 'RELEASED' else '')
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0xB4, 0x54, 0x00)
    run.font.bold = True

    doc.add_paragraph()

    # ── Employee details ────────────────────────────────────────────────
    info = doc.add_table(rows=0, cols=4)
    info.autofit = True
    pairs = [
        ('Employee Name', payslip.get('employee_name') or '—'),
        ('Employee ID', payslip.get('employee_code') or '—'),
        ('Designation', payslip.get('designation') or '—'),
        ('Department', payslip.get('department') or '—'),
    ]
    row = info.add_row().cells
    for i, (label, value) in enumerate(pairs):
        row[i].text = ''
        p1 = row[i].paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = MUTED
        p2 = row[i].add_paragraph()
        r2 = p2.add_run(value)
        r2.font.size = Pt(10.5)
        r2.font.bold = True
        _set_cell_borders(row[i])

    doc.add_paragraph()

    # ── Attendance summary ──────────────────────────────────────────────
    att_title = doc.add_paragraph()
    r = att_title.add_run('Attendance Summary')
    r.font.bold = True
    r.font.size = Pt(11)

    att = doc.add_table(rows=0, cols=4)
    row = att.add_row().cells
    for i, (label, value) in enumerate([
        ('Working Days', payslip.get('working_days')),
        ('Present', payslip.get('present_days')),
        ('On Leave', payslip.get('leave_days')),
        ('Absent (LOP)', payslip.get('absent_days')),
    ]):
        row[i].text = ''
        p1 = row[i].paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = MUTED
        p2 = row[i].add_paragraph()
        r2 = p2.add_run(str(value if value is not None else '—'))
        r2.font.size = Pt(10.5)
        r2.font.bold = True
        _set_cell_borders(row[i])

    doc.add_paragraph()

    # ── Earnings / Deductions ───────────────────────────────────────────
    cols_table = doc.add_table(rows=1, cols=2)
    cols_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left_cell, right_cell = cols_table.rows[0].cells

    for cell, title, rows, total_label, total_value in [
        (left_cell, 'Earnings', [
            ('Basic', payslip.get('basic')),
            ('HRA', payslip.get('hra')),
            ('DA', payslip.get('da')),
            ('Allowances', payslip.get('allowances')),
        ], 'Gross Salary', payslip.get('gross_salary')),
        (right_cell, 'Deductions', [
            ('Provident Fund', payslip.get('pf_deduction')),
            ('ESI', payslip.get('esi_deduction')),
            ('Income Tax', payslip.get('income_tax')),
            ('Professional Tax / Other', payslip.get('other_deductions')),
        ], 'Total Deductions', payslip.get('total_deductions')),
    ]:
        cell.text = ''
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(11)

        t = cell.add_table(rows=0, cols=2)
        for label, value in rows:
            _row(t, label, _inr(value))
        _row(t, total_label, _inr(total_value), bold=True, shade='EEF3FC')
        _set_cell_borders(cell)

    doc.add_paragraph()

    # ── Net pay ──────────────────────────────────────────────────────────
    net_table = doc.add_table(rows=0, cols=2)
    _row(net_table, 'NET PAY', _inr(payslip.get('net_salary')), bold=True, shade='E9F8F0')
    for cell in net_table.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(13)

    if payslip.get('remarks'):
        doc.add_paragraph()
        rem = doc.add_paragraph()
        r = rem.add_run(f"Remarks: {payslip['remarks']}")
        r.font.size = Pt(9.5)
        r.italic = True

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run('This is a system-generated payslip and does not require a signature.')
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    doc.save(output_path)
    return output_path
